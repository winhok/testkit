"""Synthetic regressions for the existing script CR; never touches business data/devices."""
import argparse
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]


def module(name, relative):
    path = ROOT / relative
    sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    sys.modules[name] = value
    spec.loader.exec_module(value)
    return value


api = module('oldcr_import_api', 'skills/api-test-automation/scripts/import_api.py')
workflows = module('oldcr_workflows', 'skills/api-test-automation/scripts/run_workflows.py')
sources = sys.modules['source_adapters']
excel = module('oldcr_excel', 'skills/testspec-generate/scripts/generate_excel.py')
context = module('oldcr_context', 'skills/_testspec-shared/scripts/migrate_change_context.py')
testlib = module('oldcr_testlib', 'skills/_testspec-shared/scripts/validate_testlib.py')
rebuild = sys.modules['rebuild_testlib_index']
publish = module('oldcr_publish', 'skills/testspec-publish/scripts/detect_conflicts.py')
android = module('oldcr_android', 'skills/android-static-app-reverse/scripts/reverse_android_apps.py')
artifacts = module('oldcr_artifacts', 'skills/generate-api-artifacts/scripts/generate_artifacts.py')
case_validator = module('oldcr_case_validator', 'skills/_testspec-shared/scripts/validate_testcases.py')
reconciliation = module('oldcr_reconciliation', 'skills/testspec-import/scripts/validate_reconciliation.py')
questions = sys.modules['validate_question_graph']
anchors = module('oldcr_anchors', 'skills/android-static-app-reverse/scripts/find_static_anchors.py')
run_api = module('oldcr_run_api', 'skills/api-test-automation/scripts/run_api.py')
legacy = module('oldcr_legacy', 'skills/api-test-automation/scripts/legacy_case_adapter.py')
snapshot = module('oldcr_snapshot', 'skills/testspec-code-calibrate/scripts/collect_change_snapshot.py')


class LegacyScriptCR(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def write(self, name, value):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False))
        return path

    def schema(self, name='openapi.json'):
        return self.write(name, {'openapi':'3.1.0','info':{'title':'synthetic','version':'1'},'paths':{'/health':{'get':{'responses':{'200':{'description':'ok'}}}}}})

    def case(self):
        return {'id':'CASE-1','title':'登录_凭据_成功','feature':'登录','priority':'P1','type':'冒烟',
                'status':'active','steps':'1、点击登录','expected_result':'1、显示首页','tp_refs':['TP-1'],
                'source_change':'synthetic','created_at':'2026-01-01','updated_at':'2026-01-01',
                'origin':{'kind':'testspec-native'},'trust':{'status':'verified'}}

    def feature(self, cases):
        return {'schema_version':2,'module':'登录','module_key':'LOGIN','feature':'凭据','feature_key':'CRED',
                'last_updated':'2026-01-01','case_count':len(cases),'related_features':[], 'cases':cases}

    def test_api_inspect_force_cannot_overwrite_source_or_hardlink(self):
        source = self.schema()
        alias = self.root / 'alias.json'
        os.link(source, alias)
        before = source.read_bytes()
        for target in (source, alias):
            self.assertEqual(api.main(['inspect',str(source),'--output',str(target),'--force']),2)
            self.assertEqual(source.read_bytes(), before)

    def test_api_import_force_cannot_overwrite_source(self):
        source = self.schema()
        before = source.read_bytes()
        self.assertEqual(api.main(['import',str(source),'--output-dir',str(self.root),'--description-name',source.name,'--force']),2)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root/'source-manifest.json').exists())

    def test_artifact_manifest_cannot_overwrite_input(self):
        source = self.schema('artifact-manifest.json')
        before = source.read_bytes()
        args = argparse.Namespace(source=str(source),target=['apifox'],output_dir=str(self.root),force=True)
        with self.assertRaises(artifacts.ArtifactError):
            artifacts.generate(args)
        self.assertEqual(source.read_bytes(), before)

    def test_legacy_case_import_cannot_overwrite_original(self):
        source = self.write('cases.json',[self.case()])
        before = source.read_bytes()
        result = subprocess.run([sys.executable,str(ROOT/'skills/testspec-import/scripts/import_legacy_cases.py'), '--input',str(source),'--output',str(source),'--overwrite'],capture_output=True,timeout=15)
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(source.read_bytes(), before)

    def test_staging_outputs_resolve_aliases(self):
        source = self.write('cases.json',[self.case()])
        staging = self.write('stage.json',{'keep':True})
        alias = self.root/'reconciliation.json'
        alias.symlink_to(staging)
        before = staging.read_bytes()
        result = subprocess.run([sys.executable,str(ROOT/'skills/testspec-import/scripts/import_legacy_cases.py'),'--input',str(source),'--output',str(staging),'--reconciliation-output',str(alias),'--overwrite'],capture_output=True,timeout=15)
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(staging.read_bytes(),before)

    def test_data_limit_never_silently_drops_rows(self):
        source = self.write('data.json',[{'id':1},{'id':2}])
        with self.assertRaises(workflows.CliConfigurationError):
            workflows.load_datasets(source,max_runs=1)
        self.assertEqual(len(workflows.load_datasets(source,max_runs=2)),2)

    def test_remote_error_does_not_repeat_token(self):
        with patch.object(sources,'_open_request',side_effect=RuntimeError('request token=synthetic-secret')):
            with self.assertRaises(sources.SourceError) as caught:
                sources._request_json('https://api.example.invalid/spec?token=synthetic-secret')
        self.assertNotIn('synthetic-secret',str(caught.exception))

    def test_company_excel_template_and_literal_formula_text(self):
        from openpyxl import load_workbook
        case = self.case()
        case['steps'] = '=1+1'
        output = self.root/'cases.xlsx'
        excel.create_excel_with_openpyxl([case],str(output))
        workbook = load_workbook(output,data_only=False)
        try:
            self.assertEqual(workbook.sheetnames,['测试用例','冒烟用例'])
            headers = ['编号','用例标题','级别','预置条件','操作步骤','测试预期内容','执行结果','执行人','执行日期','备注']
            for sheet in workbook:
                self.assertEqual([sheet.cell(1,i).value for i in range(1,11)], headers)
                self.assertEqual(sheet['E2'].value,'=1+1')
                self.assertEqual(sheet['E2'].data_type,'s')
                self.assertEqual(sheet.column_dimensions['E'].width,40)
                self.assertTrue(sheet['A1'].font.bold)
            self.assertEqual(workbook['测试用例']['A1'].fill.fgColor.rgb,'004472C4')
            self.assertEqual(workbook['冒烟用例']['A1'].fill.fgColor.rgb,'00548235')
        finally:
            workbook.close()

    def test_both_exports_protect_source_without_changing_format(self):
        source = self.write('cases.json',[self.case()])
        before = source.read_bytes()
        for name in ('generate_excel.py','generate_xmind.py'):
            result = subprocess.run([sys.executable,str(ROOT/'skills/testspec-generate/scripts'/name),'--input',str(source),'--output',str(source)],capture_output=True,timeout=15)
            self.assertNotEqual(result.returncode,0)
            self.assertEqual(source.read_bytes(),before)

    def test_context_migration_handles_compact_and_nested_metadata(self):
        for name, indent in [('compact.json',None),('pretty.json',2)]:
            root = {'metadata':{'_context':{'keep':'unchanged'}},'_context':{'source_skill':'old'},'testcases':[self.case()]}
            path = self.root/name
            content = json.dumps(root,indent=indent)
            path.write_text(content)
            context._write_artifact(path,content,{'context_schema_version':2})
            result = json.loads(path.read_text())
            self.assertEqual(result['_context'],{'context_schema_version':2})
            self.assertEqual(result['metadata'],root['metadata'])
            self.assertEqual(result['testcases'],root['testcases'])

    def test_duplicate_top_level_context_rejected_before_writes(self):
        path = self.root/'bad.json'
        path.write_text('{"_context":{},"_context":{}}')
        with self.assertRaises(ValueError):
            context._load_artifact(path)

    def test_duplicate_ids_in_same_testlib_file_block(self):
        self.write('modules/login/cred.json',self.feature([self.case(),self.case()]))
        rebuild.rebuild(self.root,'2026-01-01')
        result = testlib.validate(self.root,'2026-01-01')
        self.assertEqual(result['status'],'fail')
        self.assertIn('DUPLICATE_CASE_ID',{issue['type'] for issue in result['issues']})

    def test_malformed_testlib_documents_return_diagnostics(self):
        for doc in ([],self.feature([None])):
            self.write('modules/login/cred.json',doc)
            self.assertEqual(testlib.validate(self.root,'2026-01-01')['status'],'fail')

    def test_rebuild_cannot_follow_metadata_symlink(self):
        source = self.write('modules/login/cred.json',self.feature([self.case()]))
        (self.root/'index.json').symlink_to(source)
        before = source.read_bytes()
        with self.assertRaises(ValueError):
            rebuild.rebuild(self.root,'2026-01-01')
        self.assertEqual(source.read_bytes(),before)

    def test_duplicate_incoming_publish_ids_block(self):
        incoming = self.write('incoming.json',{'_context':{'origin':{'kind':'testspec-native'},'trust':{'status':'verified'}},'testcases':[self.case(),self.case()]})
        self.assertGreater(publish.detect(incoming,self.root/'empty')['hard_block_count'],0)

    def test_android_force_cannot_delete_input(self):
        output = self.root/'jadx'
        source = self.write('jadx/app.apk',{'synthetic':True})
        before = source.read_bytes()
        with self.assertRaises(RuntimeError):
            android.run_jadx([source],output,True,'auto',False)
        self.assertEqual(source.read_bytes(),before)

    def test_android_reports_failed_tool_exit_as_nonzero(self):
        source = self.write('app.apk',{'synthetic':True})
        result = android.AppResult('app','synthetic.app',self.root,self.root,None,None,None,None,None,None,None,1,2,None,'failed-exit-2')
        with patch.object(android,'require_tool'), patch.object(android,'process_app',return_value=result):
            code = android.main([str(source),'--out',str(self.root/'out')])
        self.assertEqual(code,1)

    def test_runner_timeout_produces_error_report(self):
        source = self.schema()
        output = self.root/'run.json'
        with patch.object(run_api,'_runner_executable',return_value='/synthetic/runner'), patch.object(run_api.subprocess,'run',side_effect=subprocess.TimeoutExpired('runner',1)):
            code = run_api.main([str(source),'--url','https://api.example.invalid','--output',str(output),'--runner-timeout','1'])
        result = json.loads(output.read_text())
        self.assertEqual(code,2)
        self.assertEqual(result['status'],'error')
        self.assertEqual(result['error_type'],'timeout')

    def test_runner_capture_uses_spools_and_reports_truncation(self):
        source = self.schema()
        output = self.root/'run.json'
        def run(command, **kwargs):
            self.assertIn('timeout',kwargs)
            self.assertNotIn('capture_output',kwargs)
            kwargs['stdout'].write(b'x'*100)
            return subprocess.CompletedProcess(command,0,None,None)
        with patch.object(run_api,'_runner_executable',return_value='/synthetic/runner'), patch.object(run_api.subprocess,'run',side_effect=run):
            code = run_api.main([str(source),'--url','https://api.example.invalid','--output',str(output),'--runner-output-limit','10'])
        result = json.loads(output.read_text())
        self.assertEqual(code,0)
        self.assertTrue(result['stdout_truncated'])
        self.assertEqual(result['stdout'],'x'*10+'\n[TRUNCATED]')

    def test_workflow_response_limit_and_boolean_equality(self):
        engine = sys.modules['workflow_engine']
        response = Mock(status=200,headers={})
        response.read.return_value = b'12345'
        manager = Mock()
        manager.__enter__ = Mock(return_value=response)
        manager.__exit__ = Mock(return_value=False)
        transport = engine.UrllibTransport()
        transport._opener = Mock()
        transport._opener.open.return_value = manager
        with patch.object(engine,'MAX_DOCUMENT_BYTES',4), self.assertRaises(engine.WorkflowTransportError):
            transport.request(method='GET',url='https://api.example.invalid',headers={},query={},body=None,timeout=1)
        response.read.assert_called_once_with(5)
        self.assertFalse(engine._compare(True,'==',1))
        self.assertTrue(engine._compare(True,'!=',1))

    def test_credential_redirect_and_https_downgrade_are_blocked(self):
        request = sources.urllib.request.Request('https://api.example.invalid/spec?token=synthetic-secret')
        handler = sources._CredentialSafeRedirect()
        for target in ('https://other.example.invalid/spec','http://api.example.invalid/spec'):
            with self.assertRaises(sources.urllib.error.URLError):
                handler.redirect_request(request,None,302,'Found',{},target)

    def test_import_timeout_reaches_remote_loader(self):
        payload = sources.LoadedPayload(json.loads(self.schema().read_text()),b'{}','https://api.example.invalid/spec')
        with patch.object(sources,'_request_json',return_value=payload) as request:
            sources.import_source('https://api.example.invalid/spec',timeout=3)
        self.assertEqual(request.call_args.kwargs['timeout'],3)

    def test_array_case_input_and_missing_points_are_diagnosed(self):
        source = self.write('cases.json',[self.case()])
        result = case_validator.validate(str(source))
        self.assertEqual(result['summary']['total_cases'],1)
        self.assertEqual(case_validator.validate(str(source),str(self.root/'missing.md'))['status'],'ERROR')

    def test_reconciliation_rejects_duplicate_and_empty_imports(self):
        report = self.write('reconciliation.json',{'records':[], 'summary':{}, '_context':{'canonical_source_policy':'prd-first','status':'ready-for-generate'}})
        for cases in ([],[{'id':'OLD-1'},{'id':'OLD-1'}],[None]):
            source = self.write('import.json',{'testcases':cases})
            self.assertTrue(reconciliation.validate(source,report,True))

    def test_question_graph_malformed_enums_and_deep_chain(self):
        question = {'id':'Q-001','kind':'fact','status':[], 'question':'synthetic?', 'depends_on':[], 'blocks_stages':[], 'recommendation':None,'resolution':None}
        result = questions.validate_context({'context_schema_version':2,'questions':[question]},target_stage='analysis')
        self.assertTrue(result)
        graph = {f'Q-{i:03d}':[f'Q-{i-1:03d}'] if i else [] for i in range(1200)}
        self.assertEqual(questions._cycles(graph),[])

    def test_static_anchor_secret_is_redacted_in_non_auth_group(self):
        result = anchors.redact('api_key="synthetic-secret-value"; fetch("/api");','urls')
        self.assertNotIn('synthetic-secret-value',result)

    def test_zip_path_preflight_leaves_no_partial_extraction(self):
        archive = self.root/'input.xapk'
        with zipfile.ZipFile(archive,'w') as writer:
            writer.writestr('valid.apk','synthetic')
            writer.writestr('../escaped.apk','synthetic')
        destination = self.root/'extract'
        destination.mkdir()
        with self.assertRaises(RuntimeError):
            android.safe_extract_zip(archive,destination)
        self.assertEqual(list(destination.iterdir()),[])

    def test_audit_report_cannot_replace_feature(self):
        feature = self.write('modules/login/cred.json',self.feature([self.case()]))
        rebuild.rebuild(self.root,'2026-01-01')
        before = feature.read_bytes()
        result = subprocess.run([sys.executable,str(ROOT/'skills/testspec-audit/scripts/audit_testlib.py'),'--testlib',str(self.root),'--output',str(feature),'--overwrite'],capture_output=True,timeout=15)
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(feature.read_bytes(),before)

    def test_legacy_migration_protects_nested_sources_and_rejects_empty(self):
        project = self.write('project.json',{'project':{'name':'synthetic'}})
        schema = self.schema()
        source = self.write('cases/old.json',{'cases':[]})
        before = source.read_bytes()
        with self.assertRaisesRegex(legacy.LegacyMigrationError,'overwrite'):
            legacy.migrate_legacy_project(project,schema,source)
        self.assertEqual(source.read_bytes(),before)
        with self.assertRaisesRegex(legacy.LegacyMigrationError,'No legacy'):
            legacy.migrate_legacy_project(project,schema,self.root/'new.yaml')
        self.assertFalse((self.root/'new.yaml').exists())

    def test_snapshot_scope_cannot_use_git_pathspec_magic(self):
        self.assertFalse(snapshot.safe_scope(':(top,glob)**'))
        self.assertTrue(snapshot.safe_scope('src/profile'))


if __name__ == '__main__':
    unittest.main()
